import base64
import io
import os
from typing import Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import (
    ParagraphStyle,
    getSampleStyleSheet
)
from reportlab.lib.units import mm
from reportlab.platypus import (
    Image as ReportLabImage,
    Paragraph,
    SimpleDocTemplate,
    Spacer
)


def decode_logo(
    logo_base64: Optional[str]
) -> Optional[bytes]:

    if not logo_base64:
        return None

    try:

        encoded = logo_base64.split(
            ",",
            1
        )[-1]

        data = base64.b64decode(
            encoded,
            validate=True
        )

    except Exception:

        return None

    max_size = int(
        os.getenv(
            "MAX_LOGO_BYTES",
            "2097152"
        )
    )

    if len(data) > max_size:
        return None

    return data


def make_txt(text: str) -> bytes:

    return text.encode(
        "utf-8"
    )


def make_docx(
    text: str,
    document_type: str,
    logo_base64: Optional[str] = None
) -> bytes:

    document = Document()

    section = document.sections[0]

    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    logo = decode_logo(
        logo_base64
    )

    if logo:

        try:

            paragraph = document.add_paragraph()

            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
            )

            run = paragraph.add_run()

            run.add_picture(
                io.BytesIO(logo),
                width=Inches(1.15)
            )

        except Exception:

            pass

    title = document.add_paragraph()

    title.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    title_run = title.add_run(
        document_type.upper()
    )

    title_run.bold = True
    title_run.font.name = "Times New Roman"
    title_run.font.size = Pt(16)

    for raw_line in text.splitlines():

        line = raw_line.strip()

        if not line:

            document.add_paragraph()

            continue

        paragraph = document.add_paragraph()

        paragraph.paragraph_format.space_after = Pt(6)

        is_heading = (
            line.isupper()
            or line.startswith("ARTICLE ")
            or line.startswith("SECTION ")
            or (
                len(line) > 2
                and line[0].isdigit()
                and line[1] in ".)"
            )
        )

        run = paragraph.add_run(
            line
        )

        run.font.name = "Times New Roman"
        run.font.size = Pt(11)

        if is_heading:

            run.bold = True

    footer = section.footer.paragraphs[0]

    footer.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    footer_run = footer.add_run(
        "LegalEase - AI-assisted draft. "
        "Review before signing."
    )

    footer_run.font.name = "Times New Roman"
    footer_run.font.size = Pt(8)

    buffer = io.BytesIO()

    document.save(buffer)

    return buffer.getvalue()


def make_pdf(
    text: str,
    document_type: str,
    logo_base64: Optional[str] = None
) -> bytes:

    buffer = io.BytesIO()

    document = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=18 * mm,
        leftMargin=18 * mm,
        topMargin=22 * mm,
        bottomMargin=20 * mm,
        title=document_type
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "LegalTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=16,
        leading=20,
        alignment=TA_CENTER,
        spaceAfter=14
    )

    body_style = ParagraphStyle(
        "LegalBody",
        parent=styles["BodyText"],
        fontName="Times-Roman",
        fontSize=10.5,
        leading=15,
        spaceAfter=7
    )

    heading_style = ParagraphStyle(
        "LegalHeading",
        parent=body_style,
        fontName="Times-Bold",
        spaceBefore=7,
        spaceAfter=7
    )

    story = []

    logo = decode_logo(
        logo_base64
    )

    if logo:

        try:

            image = ReportLabImage(
                io.BytesIO(logo),
                width=28 * mm,
                height=28 * mm
            )

            image.hAlign = "CENTER"

            story.append(image)

            story.append(
                Spacer(1, 5)
            )

        except Exception:

            pass

    story.append(
        Paragraph(
            document_type.upper(),
            title_style
        )
    )

    for raw_line in text.splitlines():

        line = raw_line.strip()

        if not line:

            story.append(
                Spacer(1, 4)
            )

            continue

        escaped = (
            line
            .replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
        )

        is_heading = (
            line.isupper()
            or line.startswith("ARTICLE ")
            or line.startswith("SECTION ")
            or (
                len(line) > 2
                and line[0].isdigit()
                and line[1] in ".)"
            )
        )

        style = (
            heading_style
            if is_heading
            else body_style
        )

        story.append(
            Paragraph(
                escaped,
                style
            )
        )

    def footer(
        canvas,
        doc
    ):

        canvas.saveState()

        canvas.setFont(
            "Helvetica",
            8
        )

        canvas.drawCentredString(
            A4[0] / 2,
            10 * mm,
            "LegalEase - AI-assisted draft. "
            "Review before signing."
        )

        canvas.restoreState()

    document.build(
        story,
        onFirstPage=footer,
        onLaterPages=footer
    )

    return buffer.getvalue()